import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from bookingapp.models import Booking
from profileapp.models import Profile


def slot(label):
    return label.split('(')[0], label


MORNING = [
    slot('09:20'), slot('09:25'), slot('09:30'), slot('09:45'), slot('09:50'), slot('09:55'),
    slot('10:05'), slot('10:10'), slot('10:15'), slot('10:20(연골)'), slot('10:30'), slot('10:35'),
    slot('10:40'), slot('10:50'), slot('10:55'), slot('11:00'), slot('11:05(연골)'), slot('11:15'),
    slot('11:20'), slot('11:25'), slot('11:35'), slot('11:40'), slot('11:45'), slot('11:50(연골)'),
    slot('12:00'), slot('12:05'), slot('12:10'), slot('12:20'), slot('12:25'), slot('12:30'),
]
AFTERNOON = [
    slot('14:20'), slot('14:25'), slot('14:30'), slot('14:45'), slot('14:50'), slot('14:55'),
    slot('15:05'), slot('15:10'), slot('15:15'), slot('15:20(연골)'), slot('15:30'), slot('15:35'),
    slot('15:40'), slot('15:50'), slot('15:55'), slot('16:00'), slot('16:05(연골)'), slot('16:15'),
    slot('16:20'), slot('16:25'), slot('16:35'), slot('16:40'), slot('16:45'), slot('16:50(연골)'),
    slot('17:00'), slot('17:05'), slot('17:10'), slot('17:20'), slot('17:25'), slot('17:30'),
]
NIGHT = [
    slot('17:45'), slot('17:50'), slot('17:55'), slot('18:05'), slot('18:10'), slot('18:15'),
    slot('18:30'), slot('18:35'), slot('18:40'), slot('18:50'), slot('18:55'), slot('19:00'),
    slot('19:05(연골)'), slot('19:15'), slot('19:20'), slot('19:25'), slot('19:35'), slot('19:40'),
    slot('19:45'), slot('19:50(연골)'), slot('20:00'), slot('20:05'),
]
KNOWN_TIMES = {key for key, _label in MORNING + AFTERNOON + NIGHT}

THIN = Border(
    left=Side(style='thin', color='D1D5DB'),
    right=Side(style='thin', color='D1D5DB'),
    top=Side(style='thin', color='D1D5DB'),
    bottom=Side(style='thin', color='D1D5DB'),
)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
HEADER_FILL = PatternFill('solid', fgColor='F8FAFC')
SUNDAY_FILL = PatternFill('solid', fgColor='FFF7F7')
SATURDAY_FILL = PatternFill('solid', fgColor='F7FAFF')
STATUS_FILLS = {
    'B5B5B5': PatternFill('solid', fgColor='B5B5B5'),
    'FFB6C1': PatternFill('solid', fgColor='FFB6C1'),
    'FFE8F4': PatternFill('solid', fgColor='FFE8F4'),
    'F5F5DC': PatternFill('solid', fgColor='F5F5DC'),
    'E9D5FF': PatternFill('solid', fgColor='E9D5FF'),
    'ADD8E6': PatternFill('solid', fgColor='ADD8E6'),
}


def weekday_name(day):
    return '월화수목금토일'[day.weekday()]


def profile_of(user):
    if user is None:
        return None
    try:
        return user.profile
    except Profile.DoesNotExist:
        return None


def booking_cell(booking):
    user = booking.user
    profile = profile_of(user)
    real_name = ((profile.real_name or '').strip() if profile else '')
    chart_num = ((profile.chart_num or '').strip() if profile else '')
    rn = (booking.booking_rn or '').strip()
    is_super = bool(user and user.is_superuser)
    status = booking.booking_status
    if status == '예약승인':
        if real_name:
            text = ' '.join(part for part in (real_name, chart_num, rn) if part)
            fill = 'B5B5B5' if is_super else 'FFB6C1'
        else:
            text = rn
            fill = 'B5B5B5' if is_super else 'FFE8F4'
        return text, fill
    if status == '예약요청':
        return ' '.join(part for part in (real_name, chart_num, rn) if part), 'F5F5DC'
    if status == '선택중':
        return '선택중', 'E9D5FF'
    if status == '예약가능':
        return '예약가능', 'ADD8E6'
    return '', None


def week_spans(today=None):
    if today is None:
        today = datetime.datetime.today()
    if isinstance(today, datetime.datetime):
        today = today.date()
    week1 = [today + datetime.timedelta(days=offset) for offset in range(7)]
    week2 = [today + datetime.timedelta(days=offset) for offset in range(7, 15)]
    return week1, week2


def latest_bookings(dates):
    chosen = {}
    rows = Booking.objects.filter(booking_date__in=dates).select_related('user').order_by('id')
    for booking in rows:
        chosen[(booking.booking_date, booking.booking_time)] = booking
    return chosen


def sheet_rows(dates, bookings):
    rows = [('dates', ''), ('weekdays', '')]
    rows.extend(('slot', key, label) for key, label in MORNING)
    rows.append(('dates', '점심시간'))
    rows.append(('weekdays', ''))
    rows.extend(('slot', key, label) for key, label in AFTERNOON)
    rows.append(('dates', '야간진료'))
    rows.append(('weekdays', '(수요일만)'))
    rows.extend(('slot', key, label) for key, label in NIGHT)
    date_set = set(dates)
    extras = sorted({
        booking_time
        for (day, booking_time) in bookings
        if day in date_set and booking_time not in KNOWN_TIMES
    })
    rows.extend(('slot', key, key) for key in extras)
    return rows


def apply_fill(cell, fill):
    if fill is not None:
        cell.fill = fill


def write_sheet(ws, dates, bookings):
    weekend = [weekday_name(day) for day in dates]
    for row_index, row in enumerate(sheet_rows(dates, bookings), start=1):
        kind = row[0]
        label = row[1] if kind != 'slot' else row[2]
        ws.cell(row_index, 1, label)
        ws.row_dimensions[row_index].height = 20
        for offset, day in enumerate(dates, start=2):
            cell = ws.cell(row_index, offset)
            day_name = weekend[offset - 2]
            fill = HEADER_FILL if kind in ('dates', 'weekdays') else None
            if day_name == '일':
                fill = SUNDAY_FILL
            elif day_name == '토':
                fill = SATURDAY_FILL
            if kind == 'dates':
                cell.value = day.strftime('%Y-%m-%d')
                cell.font = Font(name='Malgun Gothic', bold=True, size=9, color='C81E1E' if day_name == '일' else ('1D4ED8' if day_name == '토' else '111827'))
            elif kind == 'weekdays':
                cell.value = day_name
                cell.font = Font(name='Malgun Gothic', bold=True, size=9, color='C81E1E' if day_name == '일' else ('1D4ED8' if day_name == '토' else '111827'))
            else:
                booking = bookings.get((day, row[1]))
                text, status_fill = ('', None)
                if booking is not None:
                    text, status_fill = booking_cell(booking)
                cell.value = text or None
                cell.font = Font(name='Malgun Gothic', size=9, color='111827')
                if status_fill:
                    fill = STATUS_FILLS[status_fill]
            apply_fill(cell, fill)
            cell.alignment = CENTER
            cell.border = THIN
        label_cell = ws.cell(row_index, 1)
        label_cell.font = Font(name='Malgun Gothic', bold=kind != 'slot', size=9, color='111827')
        label_cell.alignment = CENTER
        label_cell.border = THIN
        if kind in ('dates', 'weekdays'):
            apply_fill(label_cell, HEADER_FILL)
    ws.column_dimensions['A'].width = 14
    for offset in range(2, len(dates) + 2):
        ws.column_dimensions[get_column_letter(offset)].width = 18
    ws.freeze_panes = 'B3'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '1:2'
    ws.page_setup.horizontalCentered = True
    ws.sheet_view.showGridLines = False
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.4


def build_workbook(today=None):
    week1, week2 = week_spans(today)
    bookings = latest_bookings(week1 + week2)
    workbook = Workbook()
    first = workbook.active
    first.title = '1주차'
    write_sheet(first, week1, bookings)
    second = workbook.create_sheet('2주차')
    write_sheet(second, week2, bookings)
    return workbook


def backup_bytes(today=None, now=None):
    now = now or datetime.datetime.now()
    stamp = now.strftime('%Y%m%d_%H%M%S')
    workbook = build_workbook(today or now)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), stamp, f'예약백업_{stamp}.xlsx'
